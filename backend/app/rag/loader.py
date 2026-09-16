"""文档加载（LangChain Document Loader 语义）。

支持：PDF / DOCX / TXT / Markdown / CSV / XLSX
每个加载器统一返回 `list[langchain_core.documents.Document]`，
并在 metadata 中带上页码 / 工作表 / 行号等定位信息，供 RAG 引用来源使用。

后续扩展（PPTX / HTML / URL / OCR）只需在 `_LOADERS` 注册一个函数。
"""

from __future__ import annotations

import asyncio
import csv
import hashlib
import io
import re
from collections.abc import Callable
from pathlib import Path
from typing import Any

from langchain_core.documents import Document

from app.core.exceptions import FileParseError, FileTypeUnsupportedError
from app.core.logging import get_logger

logger = get_logger(__name__)

# 扩展名 -> 内部类型
EXTENSION_MAP: dict[str, str] = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".doc": "docx",
    ".txt": "txt",
    ".text": "txt",
    ".md": "md",
    ".markdown": "md",
    ".csv": "csv",
    ".xlsx": "xlsx",
    ".xlsm": "xlsx",
}

SUPPORTED_EXTENSIONS = tuple(EXTENSION_MAP.keys())

_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_MULTI_BLANK = re.compile(r"\n{3,}")
_MULTI_SPACE = re.compile(r"[ \t\u3000]{2,}")


def detect_file_type(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    file_type = EXTENSION_MAP.get(ext)
    if not file_type:
        raise FileTypeUnsupportedError(
            f"暂不支持的文件类型 {ext or '未知'}，当前支持：{'、'.join(sorted(set(SUPPORTED_EXTENSIONS)))}"
        )
    return file_type


def compute_file_hash(path: Path, chunk_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fp:
        while True:
            block = fp.read(chunk_size)
            if not block:
                break
            digest.update(block)
    return digest.hexdigest()


def clean_text(text: str) -> str:
    """文本清洗：去掉控制字符、压缩多余空白、统一换行。"""
    if not text:
        return ""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\u00a0", " ")
    text = _CONTROL_CHARS.sub("", text)
    text = _MULTI_SPACE.sub(" ", text)
    lines = [line.strip() for line in text.split("\n")]
    text = "\n".join(lines)
    return _MULTI_BLANK.sub("\n\n", text).strip()


def decode_bytes(raw: bytes) -> str:
    """多编码兜底解码（中文文档常见 GBK / GB18030）。"""
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "big5", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


# --------------------------------------------------------------------------- #
# 各类型加载器（同步实现，最终由 to_thread 调度）
# --------------------------------------------------------------------------- #
def _load_pdf(path: Path) -> list[Document]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    documents: list[Document] = []
    for page_no, page in enumerate(reader.pages, start=1):
        try:
            text = clean_text(page.extract_text() or "")
        except Exception as exc:  # noqa: BLE001 - 单页失败不影响整体
            logger.warning("PDF 第 %s 页解析失败：%s", page_no, exc)
            continue
        if text:
            documents.append(Document(page_content=text, metadata={"page": page_no}))
    if not documents:
        raise FileParseError("PDF 未提取到任何文本（可能是扫描件，需要 OCR 支持）")
    return documents


def _load_docx(path: Path) -> list[Document]:
    import docx  # python-docx

    document = docx.Document(str(path))
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            style = (paragraph.style.name or "") if paragraph.style else ""
            prefix = "## " if style.lower().startswith("heading") else ""
            blocks.append(f"{prefix}{text}")
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    text = clean_text("\n".join(blocks))
    if not text:
        raise FileParseError("DOCX 未提取到任何文本")
    return [Document(page_content=text, metadata={"paragraphs": len(document.paragraphs)})]


def _load_text(path: Path) -> list[Document]:
    text = clean_text(decode_bytes(path.read_bytes()))
    if not text:
        raise FileParseError("文件内容为空")
    return [Document(page_content=text, metadata={})]


def _load_csv(path: Path) -> list[Document]:
    raw = decode_bytes(path.read_bytes())
    reader = csv.reader(io.StringIO(raw))
    rows = [row for row in reader if any(cell.strip() for cell in row)]
    if not rows:
        raise FileParseError("CSV 内容为空")

    header = [h.strip() for h in rows[0]]
    body = rows[1:]
    documents: list[Document] = []
    batch_size = 40
    for start in range(0, max(len(body), 1), batch_size):
        batch = body[start : start + batch_size]
        lines = [" | ".join(header)]
        for row_no, row in enumerate(batch, start=start + 2):
            pairs = [f"{header[i] if i < len(header) else f'col{i}'}: {cell}" for i, cell in enumerate(row)]
            lines.append(f"[行 {row_no}] " + " | ".join(pairs))
        documents.append(Document(page_content=clean_text("\n".join(lines)), metadata={"row_start": start + 2}))
    if not documents:
        raise FileParseError("CSV 未解析出有效数据")
    return documents


def _load_xlsx(path: Path) -> list[Document]:
    from openpyxl import load_workbook

    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    documents: list[Document] = []
    try:
        for sheet in workbook.worksheets:
            rows = [row for row in sheet.iter_rows(values_only=True) if any(cell is not None for cell in row)]
            if not rows:
                continue
            header = [str(cell).strip() if cell is not None else f"col{i}" for i, cell in enumerate(rows[0])]
            body = rows[1:]
            batch_size = 40
            for start in range(0, max(len(body), 1), batch_size):
                batch = body[start : start + batch_size]
                lines = [f"工作表：{sheet.title}", " | ".join(header)]
                for row_no, row in enumerate(batch, start=start + 2):
                    pairs = [
                        f"{header[i] if i < len(header) else f'col{i}'}: {cell}"
                        for i, cell in enumerate(row)
                        if cell is not None
                    ]
                    if pairs:
                        lines.append(f"[行 {row_no}] " + " | ".join(pairs))
                documents.append(
                    Document(
                        page_content=clean_text("\n".join(lines)),
                        metadata={"sheet": sheet.title, "row_start": start + 2},
                    )
                )
    finally:
        workbook.close()
    if not documents:
        raise FileParseError("Excel 未解析出有效数据")
    return documents


_LOADERS: dict[str, Callable[[Path], list[Document]]] = {
    "pdf": _load_pdf,
    "docx": _load_docx,
    "txt": _load_text,
    "md": _load_text,
    "csv": _load_csv,
    "xlsx": _load_xlsx,
}


# --------------------------------------------------------------------------- #
# 对外接口
# --------------------------------------------------------------------------- #
async def load_documents(path: Path, file_type: str) -> tuple[list[Document], dict[str, Any]]:
    """加载文档，返回 `(documents, 统计信息)`。

    解析在独立线程中执行，避免阻塞事件循环。
    """
    if not path.exists():
        raise FileParseError("文件不存在或已被删除")
    loader = _LOADERS.get(file_type)
    if loader is None:
        raise FileTypeUnsupportedError(f"暂不支持的文档类型：{file_type}")

    try:
        documents = await asyncio.to_thread(loader, path)
    except (FileParseError, FileTypeUnsupportedError):
        raise
    except Exception as exc:  # noqa: BLE001
        logger.warning("文档解析失败 %s: %s", path.name, exc)
        raise FileParseError(f"文档解析失败：{exc}") from exc

    total_chars = sum(len(doc.page_content) for doc in documents)
    meta: dict[str, Any] = {
        "loader": file_type,
        "parts": len(documents),
        "char_count": total_chars,
    }
    if file_type == "pdf":
        meta["pages"] = len(documents)
    logger.info("文档解析完成 %s type=%s parts=%s chars=%s", path.name, file_type, len(documents), total_chars)
    return documents, meta
